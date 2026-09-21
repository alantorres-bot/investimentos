"""Leitura dos extratos analíticos do Tesouro Direto (Nu Investimentos / B3).

O extrato analítico mostra a posição de cada aplicação no fim do mês — não traz
o rendimento do período. O rendimento de um mês é obtido comparando dois meses
seguidos:

    rendimento = valor bruto do mês
               - valor bruto do mês anterior
               - aplicações feitas no mês
               + resgates ocorridos no mês

Resgates não aparecem em valor no extrato; são deduzidos pela queda na
quantidade de títulos de uma aplicação e valorados pelo preço médio do título
entre os dois meses — por isso o mês fica marcado como estimado.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

# Colunas do extrato analítico (1-based)
COL_DATA = 1
COL_QUANTIDADE = 2
COL_VALOR_INVESTIDO = 4
COL_VALOR_BRUTO = 8
COL_VALOR_LIQUIDO = 15

_RE_DATA = re.compile(r"^\d{2}/\d{2}/\d{4}$")
_RE_PERIODO = re.compile(r"Per[ií]odo\s+(\d{2})/(\d{4})")


@dataclass
class Aplicacao:
    data: str            # dd/mm/aaaa
    quantidade: float
    investido: float
    bruto: float
    liquido: float

    @property
    def preco_unitario(self) -> float:
        return self.bruto / self.quantidade if self.quantidade else 0.0


@dataclass
class Extrato:
    arquivo: str
    titulo: str                      # ex.: "Tesouro Selic 2031"
    vencimento: str
    competencia: str                 # AAAA-MM
    aplicacoes: dict[str, Aplicacao] = field(default_factory=dict)
    total_investido: float = 0.0
    total_bruto: float = 0.0
    total_liquido: float = 0.0


@dataclass
class RendimentoMes:
    competencia: str
    bruto: float
    liquido: float
    aplicacoes_mes: float
    resgates: float
    estimado: bool
    aviso: str = ""


def _numero(valor) -> float:
    """Converte '1.234,56' (ou número) em float."""
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return 0.0


def ler_extrato(origem, nome_arquivo: str = "") -> Extrato:
    """Lê um extrato analítico (.xlsx). `origem` é um caminho ou arquivo em memória."""
    try:
        wb = openpyxl.load_workbook(origem, data_only=True, read_only=False)
    except Exception as erro:  # arquivo corrompido, formato errado etc.
        raise ValueError(f"Não foi possível abrir “{nome_arquivo}”: {erro}") from None

    ws = wb.worksheets[0]
    cabecalho = str(ws.cell(1, 1).value or "")
    if "EXTRATO ANAL" not in cabecalho.upper():
        raise ValueError(
            f"“{nome_arquivo}” não parece um extrato analítico do Tesouro Direto."
        )

    titulo = cabecalho.split("-", 1)[1].strip() if "-" in cabecalho else cabecalho.strip()
    vencimento = ""
    competencia = ""
    for linha in range(2, min(ws.max_row, 30) + 1):
        texto = str(ws.cell(linha, 1).value or "")
        if texto.upper().startswith("VENCIMENTO"):
            vencimento = texto.split(":", 1)[-1].strip()
        achado = _RE_PERIODO.search(texto)
        if achado:
            competencia = f"{achado.group(2)}-{achado.group(1)}"

    extrato = Extrato(
        arquivo=nome_arquivo or str(origem),
        titulo=titulo,
        vencimento=vencimento,
        competencia=competencia,
    )

    for linha in range(2, ws.max_row + 1):
        primeira = str(ws.cell(linha, COL_DATA).value or "").strip()
        if primeira.lower() == "total":
            extrato.total_investido = _numero(ws.cell(linha, COL_VALOR_INVESTIDO).value)
            extrato.total_bruto = _numero(ws.cell(linha, COL_VALOR_BRUTO).value)
            extrato.total_liquido = _numero(ws.cell(linha, COL_VALOR_LIQUIDO).value)
            continue
        if not _RE_DATA.match(primeira):
            continue
        extrato.aplicacoes[primeira] = Aplicacao(
            data=primeira,
            quantidade=_numero(ws.cell(linha, COL_QUANTIDADE).value),
            investido=_numero(ws.cell(linha, COL_VALOR_INVESTIDO).value),
            bruto=_numero(ws.cell(linha, COL_VALOR_BRUTO).value),
            liquido=_numero(ws.cell(linha, COL_VALOR_LIQUIDO).value),
        )

    if not extrato.competencia:
        raise ValueError(
            f"Não encontrei o período (mês/ano) dentro de “{extrato.arquivo}”."
        )
    if not extrato.aplicacoes:
        raise ValueError(f"Nenhuma aplicação encontrada em “{extrato.arquivo}”.")
    if not extrato.total_bruto:
        # extrato com uma única aplicação pode não trazer linha de total
        extrato.total_investido = sum(a.investido for a in extrato.aplicacoes.values())
        extrato.total_bruto = sum(a.bruto for a in extrato.aplicacoes.values())
        extrato.total_liquido = sum(a.liquido for a in extrato.aplicacoes.values())

    return extrato


def _competencia_da_aplicacao(data: str) -> str:
    """'03/03/2026' -> '2026-03'."""
    return f"{data[6:]}-{data[3:5]}"


def calcular_rendimentos(extratos: list[Extrato]) -> list[RendimentoMes]:
    """Calcula o rendimento de cada mês a partir dos extratos de um mesmo título."""
    extratos = sorted(extratos, key=lambda e: e.competencia)
    resultado: list[RendimentoMes] = []
    anterior: Extrato | None = None

    for extrato in extratos:
        aplicacoes_mes = sum(
            a.investido
            for data, a in extrato.aplicacoes.items()
            if _competencia_da_aplicacao(data) == extrato.competencia
        )
        resgates = 0.0
        estimado = False
        aviso = ""

        if anterior is None:
            anteriores = [
                data
                for data in extrato.aplicacoes
                if _competencia_da_aplicacao(data) < extrato.competencia
            ]
            if anteriores:
                aviso = (
                    "Primeiro mês da sequência com aplicações anteriores: o valor "
                    "seria o ganho acumulado desde o início, não o do mês. "
                    "Importe também o extrato do mês anterior."
                )
                resultado.append(
                    RendimentoMes(
                        competencia=extrato.competencia,
                        bruto=extrato.total_bruto - extrato.total_investido,
                        liquido=extrato.total_liquido - extrato.total_investido,
                        aplicacoes_mes=aplicacoes_mes,
                        resgates=0.0,
                        estimado=True,
                        aviso=aviso,
                    )
                )
                anterior = extrato
                continue
            bruto = extrato.total_bruto - aplicacoes_mes
            liquido = extrato.total_liquido - aplicacoes_mes
        else:
            for data, aplicacao_ant in anterior.aplicacoes.items():
                atual = extrato.aplicacoes.get(data)
                qtd_atual = atual.quantidade if atual else 0.0
                if qtd_atual >= aplicacao_ant.quantidade - 1e-9:
                    continue
                queda = aplicacao_ant.quantidade - qtd_atual
                precos = [aplicacao_ant.preco_unitario]
                if atual and atual.quantidade:
                    precos.append(atual.preco_unitario)
                resgates += queda * (sum(precos) / len(precos))
                estimado = True
                aviso = (
                    "Mês com resgate: o extrato não informa o valor recebido, "
                    "então o resgate foi estimado pelo preço médio do título."
                )
            bruto = extrato.total_bruto - anterior.total_bruto - aplicacoes_mes + resgates
            liquido = (
                extrato.total_liquido - anterior.total_liquido - aplicacoes_mes + resgates
            )

        resultado.append(
            RendimentoMes(
                competencia=extrato.competencia,
                bruto=round(bruto, 2),
                liquido=round(liquido, 2),
                aplicacoes_mes=round(aplicacoes_mes, 2),
                resgates=round(resgates, 2),
                estimado=estimado,
                aviso=aviso,
            )
        )
        anterior = extrato

    return resultado


def agrupar_por_titulo(extratos: list[Extrato]) -> dict[str, list[Extrato]]:
    """Agrupa os extratos por título (cada título vira um investimento)."""
    grupos: dict[str, list[Extrato]] = {}
    for extrato in extratos:
        grupos.setdefault(extrato.titulo, []).append(extrato)
    for titulo in grupos:
        grupos[titulo].sort(key=lambda e: e.competencia)
    return grupos


def competencias_duplicadas(extratos: list[Extrato]) -> list[str]:
    vistos, repetidos = set(), []
    for extrato in extratos:
        if extrato.competencia in vistos:
            repetidos.append(extrato.competencia)
        vistos.add(extrato.competencia)
    return repetidos
